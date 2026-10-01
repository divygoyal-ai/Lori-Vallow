#!/usr/bin/env python3
"""Build the Clip Check xlsx and PDF for Episode 1 (v7). Run from this folder."""
import base64
import glob
import html
import json
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import data as D

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = json.load(open(os.path.join(HERE, "shots.json")))
FONTS = glob.glob("/tmp/claude-0/-home-user-Lori-Vallow/*/scratchpad/fonts/node_modules/@fontsource")[0]
EPISODE = "Lori Vallow, Episode 1, The Man Who Was Afraid (v7)"
OWN = "Own AI and graphics"


def tc(t):
    return "%d:%04.1f" % (int(t // 60), t % 60)


rows = []
for s in SHOTS:
    c = D.CLS[s["src"]]
    rows.append(dict(id=s["id"], t=s["t"], d=s["d"], src=s["src"], tag=s.get("ai") or s.get("tag") or "",
                     group=c[0], owner=c[1], how=c[2], what=c[3], use=c[4], type=c[5], verdict=c[6],
                     fix=c[7], issues=c[8]))
name, t0, d0 = D.END_CARD
rows.append(dict(id="end", t=t0, d=d0, src="end card", tag="", group=OWN, owner="Osira", how="Made by us",
                 what="End card, EPISODE 2", use="Own graphic", type="V18", verdict="USE",
                 fix="Add the credits for the Creative Commons images here.", issues="E"))

third = [r for r in rows if r["group"] != OWN]
tp_sec = sum(r["d"] for r in third)
covered = sum(r["d"] for r in rows)
VERDICTS = ["USE", "USE IF", "LICENSE", "DON'T", "NOT CLEARED"]
cnt = {v: sum(1 for r in rows if r["verdict"] == v) for v in VERDICTS}
secs = {v: round(sum(r["d"] for r in rows if r["verdict"] == v), 2) for v in VERDICTS}
STATUS = "BLOCKED"

# ---------------------------------------------------------------- xlsx
wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="14161A")
hdr_font = Font(bold=True, color="FFFFFF")


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = hdr_fill, hdr_font
    ws.freeze_panes = "A2"


def wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


ws = wb.active
ws.title = "Clip log"
header(ws, ["Shot", "Start (s)", "End (s)", "Duration (s)", "On-screen tag", "Source file", "Source group", "Original owner",
            "How obtained", "What is on screen", "Use class", "Clip type", "Verdict", "Fix", "Issues"],
       [7, 9, 9, 11, 14, 24, 28, 34, 40, 38, 24, 12, 13, 60, 8])
for i, r in enumerate(rows, 2):
    ws.append([r["id"], r["t"], round(r["t"] + r["d"], 2), "=C%d-B%d" % (i, i), r["tag"], r["src"], r["group"], r["owner"],
               r["how"], r["what"], r["use"], r["type"], r["verdict"], r["fix"], r["issues"]])
wrap(ws)
last = len(rows) + 1

bs = wb.create_sheet("By source")
header(bs, ["Source group", "Shots", "Seconds", "Share of runtime"], [40, 10, 12, 16])
groups = list(OrderedDict.fromkeys(r["group"] for r in rows))
for i, g in enumerate(groups, 2):
    bs.append([g, '=COUNTIF(\'Clip log\'!$G$2:$G$%d,A%d)' % (last, i),
               '=SUMIFS(\'Clip log\'!$D$2:$D$%d,\'Clip log\'!$G$2:$G$%d,A%d)' % (last, last, i),
               "=C%d/$B$%d" % (i, len(groups) + 5)])
n = len(groups) + 2
bs.append(["Total", "=SUM(B2:B%d)" % (n - 1), "=SUM(C2:C%d)" % (n - 1), "=C%d/$B$%d" % (n, len(groups) + 5)])
bs.append([])
bs.append(["Third-party seconds (everything except own AI and graphics)", "", "=C%d-SUMIFS(C2:C%d,A2:A%d,\"%s\")" % (n, n - 1, n - 1, OWN), ""])
bs.append(["Runtime (s)", D.TOTAL])
bs.append([])
bs.append(["Verdict", "Shots", "Seconds", "Share of runtime"])
base = n + 6
for j, v in enumerate(VERDICTS):
    r = base + j
    bs.append([v, '=COUNTIF(\'Clip log\'!$M$2:$M$%d,A%d)' % (last, r), '=SUMIFS(\'Clip log\'!$D$2:$D$%d,\'Clip log\'!$M$2:$M$%d,A%d)' % (last, last, r),
               "=C%d/$B$%d" % (r, len(groups) + 5)])
# fix: share formulas reference runtime cell; place runtime at row len(groups)+5
for i in range(2, n + 1):
    bs.cell(row=i, column=4).number_format = "0.0%"
for j in range(len(VERDICTS)):
    bs.cell(row=base + j, column=4).number_format = "0.0%"
wrap(bs)

iss = wb.create_sheet("Issues")
header(iss, ["ID", "Timecode", "Issue", "What we found", "Why it matters", "Severity", "Fix"], [6, 30, 36, 80, 50, 10, 70])
for it in D.ISSUES:
    iss.append(list(it))
wrap(iss)

q = wb.create_sheet("Questions for counsel")
header(q, ["No.", "Question"], [6, 140])
for i, t in enumerate(D.COUNSEL, 1):
    q.append([i, t])
wrap(q)

rs = wb.create_sheet("Research sources")
header(rs, ["Source", "Link or file", "Status"], [60, 80, 80])
for s in D.SOURCES:
    rs.append(list(s))
nc = wb.create_sheet("Not checked")
header(nc, ["Item"], [150])
for t in D.NOT_CHECKED:
    nc.append([t])
wrap(rs)
wrap(nc)
XLSX = os.path.join(HERE, "LoriVallow_EP01_ClipCheck_ClearanceLog.xlsx")
wb.save(XLSX)

# ---------------------------------------------------------------- pdf
def font_face(fam, pkg, weight, style="normal"):
    p = glob.glob("%s/%s/files/%s-latin-%d-%s.woff2" % (FONTS, pkg, pkg, weight, style))
    if not p:
        return ""
    b = base64.b64encode(open(p[0], "rb").read()).decode()
    return "@font-face{font-family:'%s';font-weight:%d;font-style:%s;src:url(data:font/woff2;base64,%s) format('woff2');}" % (fam, weight, style, b)


css_fonts = "".join([font_face("Oswald", "oswald", 500), font_face("Oswald", "oswald", 700),
                     font_face("Inter", "inter", 400), font_face("Inter", "inter", 600),
                     font_face("Source Serif 4", "source-serif-4", 400), font_face("Source Serif 4", "source-serif-4", 400, "italic"),
                     font_face("JetBrains Mono", "jetbrains-mono", 400), font_face("JetBrains Mono", "jetbrains-mono", 700)])
COL = {"USE": "#2F7D4F", "USE IF": "#2B5FA8", "LICENSE": "#A86B12", "DON'T": "#B3122B", "NOT CLEARED": "#A86B12"}
SEV = {"Red": "#B3122B", "Fix": "#A86B12", "Edge": "#2B5FA8"}


def chip(v, colors=COL):
    return '<span class="chip" style="background:%s">%s</span>' % (colors[v], html.escape(v))


def img(p):
    return "data:image/jpeg;base64," + base64.b64encode(open(os.path.join(HERE, p), "rb").read()).decode()


E = html.escape
STILLS = {"A": ["stills/f_127.0.jpg", "stills/f_129.5.jpg"], "B": ["stills/f_33.0.jpg", "stills/lc_src.jpg"],
          "C": ["stills/f_1.0.jpg", "stills/f_16.0.jpg", "stills/f_26.5.jpg", "stills/f_88.8.jpg"],
          "D": ["stills/f_11.0.jpg", "stills/f_66.0.jpg"], "E": ["stills/f_6.0.jpg", "stills/f_138.0.jpg"],
          "F": ["stills/f_96.5.jpg"], "G": ["stills/f_121.0.jpg"], "H": ["stills/f_66.0.jpg"]}
CAPS = {"stills/lc_src.jpg": "Source file: Law&Crime logo bottom left (not in our cut)"}

top3 = [
    "Fix the quote on screen and in the voiceover: the 'exact words' are not in the court papers (issue A).",
    "Replace the bodycam: it is a Law&Crime copy with the logo cropped out and the wrong agency tag (issue B).",
    "Get written family consent for the four photos of Charles, or replace them (issue C).",
]

h = ["<html><head><meta charset='utf-8'><style>", css_fonts, """
@page{size:A4;margin:16mm 14mm 18mm 14mm}
*{box-sizing:border-box}
body{font-family:'Inter',sans-serif;color:#14161A;font-size:9.5pt;line-height:1.45;margin:0}
h1,h2,h3{font-family:'Oswald',sans-serif;font-weight:700;margin:0;letter-spacing:.02em;text-transform:uppercase}
h2{font-size:15pt;border-bottom:1.5px solid #14161A;padding-bottom:3px;margin:18px 0 8px}
h3{font-size:11.5pt;margin:0 0 4px}
.lead{font-family:'Source Serif 4',serif;font-size:12pt}
.mono{font-family:'JetBrains Mono',monospace;font-size:8pt}
.status{font-family:'Oswald';font-weight:700;font-size:54pt;line-height:1;color:#B3122B;margin:6px 0}
.chip{display:inline-block;color:#fff;font-family:'JetBrains Mono',monospace;font-weight:700;font-size:7pt;padding:2px 6px;border-radius:2px;white-space:nowrap}
.cards{display:flex;gap:8px;margin:10px 0}
.card{flex:1;border:1px solid #E4E0D6;padding:8px;text-align:center}
.card b{font-family:'Oswald';font-size:20pt;display:block}
table{width:100%;border-collapse:collapse;font-size:7.6pt}
th{background:#14161A;color:#fff;text-align:left;padding:4px;font-family:'JetBrains Mono';font-weight:400;font-size:7pt}
td{border-bottom:1px solid #E4E0D6;padding:4px;vertical-align:top}
tr{page-break-inside:avoid}
.issue{border:1px solid #E4E0D6;padding:10px;margin:10px 0;page-break-inside:avoid}
.stills{display:flex;gap:6px;margin:6px 0;flex-wrap:wrap}
.stills figure{margin:0;width:31%}
.stills img{width:100%;border:1px solid #E4E0D6}
.stills figcaption{font-family:'JetBrains Mono';font-size:6.5pt;color:#555}
ul{margin:4px 0 4px 16px;padding:0}
.note{border-left:3px solid #A86B12;background:#faf6ec;padding:6px 10px;margin:10px 0;font-size:8.5pt}
""", "</style></head><body>"]

h.append("<div class='mono'>OSIRA · CLIP CHECK · %s</div>" % E(EPISODE))
h.append("<div class='status'>%s</div>" % STATUS)
h.append("<p class='lead'>Three red lines are still in the cut: a misquote, bodycam with the source logo cropped out and the wrong agency named, and photos of a homicide victim with no consent.</p>")
h.append("<div class='cards'>")
for v in VERDICTS:
    h.append("<div class='card'>%s<b>%d</b><span class='mono'>%.1f s</span></div>" % (chip(v), cnt[v], secs[v]))
h.append("</div>")
h.append("<p><b>Third-party material:</b> %.1f s of %.1f s (%.0f%% of runtime). Own AI and graphics: %.1f s. Video file: <span class='mono'>%s</span>, 1080x1920.</p>" % (
    tp_sec, D.TOTAL, 100 * tp_sec / D.TOTAL, D.TOTAL - tp_sec, E(D.VIDEO)))
h.append("<h2>Top three fixes</h2><ol>%s</ol>" % "".join("<li>%s</li>" % E(t) for t in top3))
h.append("<div class='note'>This is a self-check, not a sign-off. Verdicts are internal working rules under the Osira Sourcing Guide, not legal advice. A producer must approve before release; edge cases go to counsel. Most likely claimants: Law&amp;Crime (bodycam copy) and the Vallow family (Charles's photos).</div>")
h.append("<p>What is fine: the court-record pages are public record and are read aloud; the interview shots are the subject's own statements; the own AI shots show no faces. Those still need the fixes listed in the table.</p>")

h.append("<h2>Clip log, third-party shots</h2><table><tr><th>Time</th><th>Source file</th><th>What it is</th><th>Original owner</th><th>Type</th><th>Verdict</th><th>Fix</th></tr>")
for r in third:
    h.append("<tr><td class='mono'>%s<br>%.1f s</td><td class='mono'>%s</td><td>%s</td><td>%s</td><td class='mono'>%s</td><td>%s</td><td>%s</td></tr>" % (
        tc(r["t"]), r["d"], E(r["src"]), E(r["what"]), E(r["owner"]), E(r["type"]), chip(r["verdict"]), E(r["fix"])))
h.append("</table>")

h.append("<h2>Issues</h2>")
for (iid, when, title, found, why, sev, fix) in D.ISSUES:
    h.append("<div class='issue'><h3>%s. %s</h3><div class='mono'>%s &nbsp; %s</div>" % (iid, E(title), E(when), chip(sev, SEV)))
    pics = STILLS.get(iid, [])
    if pics:
        h.append("<div class='stills'>%s</div>" % "".join(
            "<figure><img src='%s'><figcaption>%s</figcaption></figure>" % (img(p), E(CAPS.get(p, "Frame from v7: " + os.path.basename(p)[2:-4] + " s"))) for p in pics))
    h.append("<p><b>Found:</b> %s</p><p><b>Why it matters:</b> %s</p><p><b>Fix:</b> %s</p></div>" % (E(found), E(why), E(fix)))

h.append("<h2>Questions for counsel</h2><ol>%s</ol>" % "".join("<li>%s</li>" % E(t) for t in D.COUNSEL))
h.append("<h2>Not checked</h2><ul>%s</ul>" % "".join("<li>%s</li>" % E(t) for t in D.NOT_CHECKED))
h.append("<h2>Sources used</h2><ul>%s</ul>" % "".join(
    "<li>%s <span class='mono'>%s</span><br><i>%s</i></li>" % (E(a), E(b), E(c)) for a, b, c in D.SOURCES))
h.append("</body></html>")
HTML = os.path.join(HERE, "report.html")
open(HTML, "w").write("".join(h))

from playwright.sync_api import sync_playwright

PDF = os.path.join(HERE, "LoriVallow_EP01_ClipCheck_Report.pdf")
footer = ("<div style=\"font-family:Inter;font-size:7px;color:#666;width:100%;padding:0 14mm;display:flex;justify-content:space-between\">"
          "<span>Osira · Clip Check · Lori Vallow EP01</span><span>Page <span class='pageNumber'></span> / <span class='totalPages'></span></span></div>")
with sync_playwright() as p:
    exe = glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome")
    b = p.chromium.launch(executable_path=exe[0]) if exe else p.chromium.launch()
    pg = b.new_page()
    pg.goto("file://" + HTML)
    pg.pdf(path=PDF, format="A4", print_background=True, display_header_footer=True, header_template="<span></span>",
           footer_template=footer, margin=dict(top="16mm", bottom="18mm", left="14mm", right="14mm"))
    b.close()
os.remove(HTML)
print("third-party %.2f s of %.1f; covered %.2f s" % (tp_sec, D.TOTAL, covered))
print(cnt, secs)
